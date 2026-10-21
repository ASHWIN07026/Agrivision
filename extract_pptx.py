import collections
import collections.abc
import pptx

prs = pptx.Presentation('Crop-Project-Presentation.pptx')
with open('presentation_text.txt', 'w', encoding='utf-8') as f:
    for i, slide in enumerate(prs.slides):
        f.write(f"--- Slide {i+1} ---\n")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                f.write(shape.text + "\n")
